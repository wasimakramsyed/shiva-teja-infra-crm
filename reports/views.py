from django.shortcuts import render
from django.http import HttpResponse
import csv
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from openpyxl import Workbook
from accounts.decorators import role_required
from reports.services.report_service import ReportService

from payments.models import Payment
from commissions.models import CommissionRecord
from bookings.models import Booking
from projects.models import Plot
from accounts.decorators import role_required


@role_required(['admin', 'manager'])
def reports_home(request):
    context = {
        'total_sales': Booking.objects.count(),

        'total_revenue': sum(
            payment.amount
            for payment in Payment.objects.all()
        ),

        'pending_commissions': CommissionRecord.objects.filter(
            status='generated'
        ).count(),
        
        'booked_plots': Plot.objects.filter(
            status='booked'
        ).count(),

        'registered_plots': Plot.objects.filter(
            status='registered'
        ).count(),
    }

    return render(
        request,
        'reports/home.html',
        context
    )


@role_required(["admin", "accounts", "manager"])
def booking_report(request):

    return render(
        request,
        "reports/booking_report.html",
    )


@role_required(["admin", "accounts", "manager"])
def payment_report(request):

    return render(
        request,
        "reports/payment_report.html",
    )


@role_required(["admin", "accounts", "manager"])
def registration_report(request):

    return render(
        request,
        "reports/registration_report.html",
    )




@role_required(["admin", "accounts", "manager"])
def commission_report(request):

    context = ReportService.commission_summary(request)

    return render(

        request,

        "reports/commission_report.html",

        context,

    )


@role_required(["admin", "accounts", "manager"])
def customer_report(request):

    return render(
        request,
        "reports/customer_report.html",
    )


@role_required(["admin", "accounts", "manager"])
def employee_report(request):

    return render(
        request,
        "reports/employee_report.html",
    )


@role_required(["admin", "accounts", "manager"])
def project_report(request):

    return render(
        request,
        "reports/project_report.html",
    )
# EXPORT BOOKINGS CSV
@role_required(['admin', 'manager'])
def export_bookings_csv(request):
    response = HttpResponse(
        content_type='text/csv'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename="bookings.csv"'

    writer = csv.writer(response)

    writer.writerow([
        'Booking ID',
        'Customer Name',
        'Project',
        'Plot',
        'Amount',
        'Status'
    ])

    bookings = Booking.objects.all()

    for booking in bookings:
        writer.writerow([
            booking.booking_id,
            booking.customer.customer_name,
            booking.project.project_name,
            booking.plot.plot_number,
            booking.booking_amount,
            booking.status
        ])

    return response


# EXPORT PAYMENTS CSV
@role_required(['admin', 'manager'])
def export_payments_csv(request):
    response = HttpResponse(
        content_type='text/csv'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename="payments.csv"'

    writer = csv.writer(response)

    writer.writerow([
        'Payment ID',
        'Booking ID',
        'Amount',
        'Payment Mode',
        'Status',
        'Payment Date'
    ])

    payments = Payment.objects.all()

    for payment in payments:
        writer.writerow([
            payment.payment_id,
            payment.booking.booking_id,
            payment.amount,
            payment.payment_mode,
            payment.status,
            payment.payment_date
        ])

    return response


# EXPORT BOOKINGS PDF
@role_required(['admin', 'manager'])
def export_bookings_pdf(request):
    response = HttpResponse(
        content_type='application/pdf'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename="bookings.pdf"'

    doc = SimpleDocTemplate(response)

    data = [[
        'Booking ID',
        'Customer',
        'Project',
        'Plot',
        'Amount',
        'Status'
    ]]

    bookings = Booking.objects.all()

    for booking in bookings:
        data.append([
            booking.booking_id,
            booking.customer.customer_name,
            booking.project.project_name,
            booking.plot.plot_number,
            str(booking.booking_amount),
            booking.status
        ])

    table = Table(data)

    elements = [table]
    doc.build(elements)

    return response


@role_required(["admin", "accounts", "manager"])
def export_commission_excel(request):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Commission Report"

    headers = [

        "Commission ID",
        "Booking",
        "Customer",
        "Project",
        "Gross",
        "TDS",
        "Net",
        "Status",

    ]

    for col, header in enumerate(headers, 1):

        sheet.cell(row=1, column=col).value = header

    commissions = CommissionRecord.objects.select_related(

        "booking",
        "customer",
        "project",

    )

    row = 2

    for commission in commissions:

        sheet.cell(row=row, column=1).value = commission.record_id

        sheet.cell(row=row, column=2).value = commission.booking.booking_id

        sheet.cell(row=row, column=3).value = commission.customer.customer_name

        sheet.cell(row=row, column=4).value = commission.project.project_name

        sheet.cell(row=row, column=5).value = float(
            commission.gross_commission
        )

        sheet.cell(row=row, column=6).value = float(
            commission.tds_amount
        )

        sheet.cell(row=row, column=7).value = float(
            commission.net_commission
        )

        sheet.cell(row=row, column=8).value = commission.status

        row += 1

    response = HttpResponse(

        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )

    response["Content-Disposition"] = (

        'attachment; filename="commission_report.xlsx"'

    )

    workbook.save(response)

    return response


@role_required(["admin", "accounts", "manager"])
def export_commission_pdf(request):

    buffer = BytesIO()

    doc = SimpleDocTemplate(buffer)

    data = [[
        "Commission ID",
        "Booking",
        "Customer",
        "Project",
        "Gross",
        "TDS",
        "Net",
        "Status",
    ]]

    commissions = CommissionRecord.objects.select_related(
        "booking",
        "customer",
        "project",
    )

    for commission in commissions:

        data.append([
            commission.record_id,
            commission.booking.booking_id,
            commission.customer.customer_name,
            commission.project.project_name,
            str(commission.gross_commission),
            str(commission.tds_amount),
            str(commission.net_commission),
            commission.status,
        ])

    table = Table(data)

    table.setStyle(TableStyle([

        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),

        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

        ("GRID", (0, 0), (-1, -1), 1, colors.black),

        ("BACKGROUND", (0, 1), (-1, -1), colors.beige),

        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),

    ]))

    doc.build([table])

    pdf = buffer.getvalue()

    buffer.close()

    response = HttpResponse(
        pdf,
        content_type="application/pdf"
    )

    response[
        "Content-Disposition"
    ] = 'attachment; filename="commission_report.pdf"'

    return response